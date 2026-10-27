# -*- coding: utf-8 -*-
"""خروجی Excel (xlsx) و CSV — برای ریز هزینه‌ها، لیست متریال و برش‌کاری."""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, Dict, List, Optional

from . import persian as fa
from .costing import Material, PageResult, ProjectResult, UnitRow

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:                            # pragma: no cover
    HAS_OPENPYXL = False

ACCENT_HEX = "0F766E"
SOFT_HEX = "E6F4F2"


# ---------------------------------------------------------------------------
def _style_header(ws, row: int, ncols: int) -> None:
    if not HAS_OPENPYXL:
        return
    fill = PatternFill("solid", fgColor=ACCENT_HEX)
    font = Font(name="Vazirmatn", bold=True, color="FFFFFF", size=11)
    align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="C9D3DE")
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)


def _autosize(ws, widths: List[int]) -> None:
    if not HAS_OPENPYXL:
        return
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _rtl(ws) -> None:
    try:
        ws.sheet_view.rightToLeft = True
    except Exception:                          # noqa: BLE001
        pass


# ---------------------------------------------------------------------------
# ۱) خروجی ریز هزینه‌های پروژه
# ---------------------------------------------------------------------------

def export_project_xlsx(result: ProjectResult, project: Dict[str, Any],
                        path: Path, settings: Optional[Dict[str, str]] = None) -> Path:
    """خروجی Excel کامل: خلاصه + ریز اقلام + آمار هر بخش + برش‌کاری."""
    if not HAS_OPENPYXL:
        raise RuntimeError("برای خروجی اکسل باید بسته‌ی openpyxl نصب باشد:  pip install openpyxl")
    settings = settings or {}
    cur = settings.get("currency", "ریال")
    scale = float(settings.get("currency_scale", "1") or 1)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # --- برگه ۱: خلاصه ---
    ws = wb.active
    ws.title = "خلاصه"
    _rtl(ws)
    ws.append([settings.get("company_name", "")])
    ws["A1"].font = Font(name="Vazirmatn", bold=True, size=15, color=ACCENT_HEX)
    ws.append([f"پروژه: {project.get('title', '')}   |   کد: {project.get('code', '')}   |   مشتری: {project.get('customer_name', '')}"])
    ws.append([f"تاریخ: {fa.jdate_long()}"])
    ws.append([])
    ws.append(["شرح", "مبلغ"])
    _style_header(ws, 5, 2)
    rows = [
        ("جمع هزینه‌ی بخش‌ها", result.pages_cost),
        ("هزینه‌ی نصب", result.installation_cost),
        ("جمع هزینه‌ی تمام‌شده", result.total_cost),
        (f"سود ({'مارکاپ' if result.profit_mode == 'markup' else 'مارجین'} {result.profit_percent:g}٪)",
         result.profit_amount),
        ("قیمت ناخالص", result.gross_price),
        ("تخفیف", -result.discount_amount),
        (f"مالیات بر ارزش افزوده ({result.vat_percent:g}٪)", result.vat_amount),
        ("مبلغ قابل پرداخت", result.final_price),
    ]
    for label, val in rows:
        ws.append([label, round(val / scale, 2)])
    for r in range(6, 6 + len(rows)):
        ws.cell(row=r, column=2).number_format = "#,##0"
        ws.cell(row=r, column=1).font = Font(name="Vazirmatn", size=11)
    ws.cell(row=5 + len(rows), column=1).font = Font(name="Vazirmatn", bold=True, size=12, color=ACCENT_HEX)
    ws.cell(row=5 + len(rows), column=2).font = Font(name="Vazirmatn", bold=True, size=12, color=ACCENT_HEX)
    ws.append([])
    ws.append([f"واحد پول: {cur}"])
    ws.append(["مبلغ به حروف: " + fa.money_to_words(result.final_price, cur, scale)])
    _autosize(ws, [46, 20])

    # --- برگه ۲: ریز اقلام ---
    ws2 = wb.create_sheet("ریز اقلام")
    _rtl(ws2)
    ws2.append(["ردیف", "گروه", "عنوان", "توضیحات", "تعداد", "واحد", "فی", "مبلغ کل", "نمایش در فاکتور"])
    _style_header(ws2, 1, 9)
    for i, it in enumerate(result.items, start=1):
        ws2.append([i, it.group, it.title, it.detail, round(it.qty, 3), it.unit_label,
                    round(it.unit_price / scale, 2), round(it.amount / scale, 2),
                    "بله" if it.visible else "خیر"])
        for c in (7, 8):
            ws2.cell(row=i + 1, column=c).number_format = "#,##0"
        for c in range(1, 10):
            cell = ws2.cell(row=i + 1, column=c)
            cell.font = Font(name="Vazirmatn", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 4))
    ws2.append([])
    ws2.append(["", "", "", "جمع کل", "", "", "", round(result.final_price / scale, 2), ""])
    _autosize(ws2, [6, 18, 30, 52, 10, 10, 14, 16, 14])
    ws2.freeze_panes = "A2"

    # --- برگه ۳: آمار بخش‌ها ---
    ws3 = wb.create_sheet("آمار بخش‌ها")
    _rtl(ws3)
    ws3.append(["بخش (سبک)", "یونیت", "لنگه در", "کشو", "متر طول زمینی", "متر طول دیواری",
                "م² ورق بدنه", "م² فیبر", "م² در", "متر نوار PVC", "ورق مصرفی", "هزینه‌ی تمام‌شده"])
    _style_header(ws3, 1, 12)
    for pr in result.pages:
        ws3.append([pr.name, pr.total_units, pr.total_doors, pr.total_drawers,
                    round(pr.total_ground_lm, 2), round(pr.total_wall_lm, 2),
                    round(pr.total_body_area, 3), round(pr.total_back_area, 3),
                    round(pr.total_door_area, 3), round(pr.total_banding_m, 2),
                    round(pr.total_sheets, 0), round(pr.total_cost / scale, 2)])
    _autosize(ws3, [22, 9, 10, 8, 15, 15, 14, 12, 12, 15, 13, 18])

    # --- برگه ۴: نقشه برش (Cut list) ---
    ws4 = wb.create_sheet("نقشه برش")
    _rtl(ws4)
    ws4.append(["بخش", "یونیت", "قطعه", "تعداد", "عرض (سانتی‌متر)", "ارتفاع (سانتی‌متر)",
                "متریال", "مساحت کل (م²)", "نوار می‌خورد"])
    _style_header(ws4, 1, 9)
    from .costing import CostingEngine
    engine = CostingEngine([], settings)
    for pr, page_data in zip(result.pages, _pages_of_result(result)):
        for u in page_data:
            for panel in engine.panels_of(u):
                if not panel.enabled:
                    continue
                qty = panel.qty * max(1, int(u.qty or 1))
                ws4.append([pr.name, u.label, panel.label, qty, panel.w_cm, panel.h_cm,
                            "فیبر پشت" if panel.material_key == "back" else "ورق بدنه",
                            round(panel.area_sqm * max(1, int(u.qty or 1)), 3),
                            "بله" if panel.banding else "خیر"])
    _autosize(ws4, [18, 20, 26, 8, 16, 18, 14, 16, 12])
    ws4.freeze_panes = "A2"

    wb.save(str(path))
    return path


def _pages_of_result(result: ProjectResult) -> List[List[UnitRow]]:
    """برای برگه‌ی نقشه برش؛ اگر داده‌ی یونیت‌ها در دسترس نبود لیست خالی برمی‌گرداند."""
    out: List[List[UnitRow]] = []
    for pr in result.pages:
        out.append(getattr(pr, "_units", []) or [])
    return out


def attach_units(result: ProjectResult, pages: List[Any]) -> None:
    """یونیت‌های هر صفحه را به نتیجه متصل می‌کند تا «نقشه برش» قابل تولید باشد."""
    for pr, page in zip(result.pages, pages):
        setattr(pr, "_units", list(getattr(page, "units", [])))


# ---------------------------------------------------------------------------
# ۲) خروجی لیست متریال
# ---------------------------------------------------------------------------

def export_materials_xlsx(materials: List[Material], path: Path,
                          cat_labels: Optional[Dict[str, str]] = None) -> Path:
    if not HAS_OPENPYXL:
        raise RuntimeError("بسته‌ی openpyxl نصب نیست.")
    cat_labels = cat_labels or {}
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "متریال و خدمات"
    _rtl(ws)
    ws.append(["شناسه", "دسته", "نام", "واحد", "قیمت", "عرض ورق (سانتی‌متر)",
               "ارتفاع ورق (سانتی‌متر)", "ضخامت (میل)", "عرض مؤثر (سانتی‌متر)",
               "قیمت هر م²", "یادداشت", "فعال"])
    _style_header(ws, 1, 12)
    for m in materials:
        ws.append([m.id, cat_labels.get(m.category, m.category), m.name, m.unit,
                   round(float(m.price), 2), m.sheet_width_cm or "", m.sheet_height_cm or "",
                   m.thickness_mm or "", m.width_cm or "",
                   round(float(m.price_per_sqm), 2) if m.price_per_sqm else "",
                   m.note or "", "بله"])
        r = ws.max_row
        ws.cell(row=r, column=5).number_format = "#,##0"
        for c in range(1, 13):
            ws.cell(row=r, column=c).font = Font(name="Vazirmatn", size=10)
    _autosize(ws, [8, 20, 36, 10, 16, 16, 18, 14, 18, 14, 34, 8])
    ws.freeze_panes = "A2"
    wb.save(str(path))
    return path


def export_materials_csv(materials: List[Material], path: Path,
                         cat_labels: Optional[Dict[str, str]] = None) -> Path:
    cat_labels = cat_labels or {}
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        w.writerow(["شناسه", "دسته", "نام", "واحد", "قیمت", "عرض ورق", "ارتفاع ورق",
                    "ضخامت", "عرض مؤثر", "یادداشت"])
        for m in materials:
            w.writerow([m.id, cat_labels.get(m.category, m.category), m.name, m.unit,
                        m.price, m.sheet_width_cm, m.sheet_height_cm, m.thickness_mm,
                        m.width_cm, m.note])
    return path


# ---------------------------------------------------------------------------
# ۳) خروجی فاکتور
# ---------------------------------------------------------------------------

def export_invoice_xlsx(data: Dict[str, Any], path: Path) -> Path:
    if not HAS_OPENPYXL:
        raise RuntimeError("بسته‌ی openpyxl نصب نیست.")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    cur = data.get("currency", "ریال")
    scale = float(data.get("scale", 1) or 1)
    wb = Workbook()
    ws = wb.active
    ws.title = data.get("kind_label", "پیش‌فاکتور")
    _rtl(ws)
    seller = data.get("seller", {})
    ws.append([seller.get("name", "")])
    ws["A1"].font = Font(name="Vazirmatn", bold=True, size=15, color=ACCENT_HEX)
    ws.append([f"{data.get('kind_label', 'پیش‌فاکتور')} شماره {data.get('number', '')} — تاریخ {data.get('date_fa', '')}"])
    cust = data.get("customer", {})
    ws.append([f"مشتری: {cust.get('name', '')}   تلفن: {cust.get('phone', '') or cust.get('mobile', '')}"])
    ws.append([f"پروژه: {data.get('project_title', '')} ({data.get('project_code', '')})"])
    ws.append([])
    ws.append(["ردیف", "شرح کالا / خدمات", "توضیحات", "تعداد", "واحد", f"فی ({cur})", f"مبلغ کل ({cur})"])
    _style_header(ws, 6, 7)
    r = 7
    for i, it in enumerate(data.get("items", []), start=1):
        ws.append([i, it.get("title", ""), it.get("detail", ""), round(it.get("qty", 0) / 1, 3),
                   it.get("unit", ""), round(it.get("unit_price", 0) / scale, 2),
                   round(it.get("amount", 0) / scale, 2)])
        for c in range(1, 8):
            cell = ws.cell(row=r, column=c)
            cell.font = Font(name="Vazirmatn", size=10)
            cell.alignment = Alignment(vertical="center", wrap_text=(c == 3))
        ws.cell(row=r, column=6).number_format = "#,##0"
        ws.cell(row=r, column=7).number_format = "#,##0"
        r += 1
    totals = data.get("totals", {})
    r += 1
    for label, key in (("جمع هزینه‌ی تمام‌شده", "total_cost"),
                       ("سود", "profit_amount"),
                       ("قیمت ناخالص", "gross_price"),
                       ("تخفیف", "discount_amount"),
                       ("مالیات بر ارزش افزوده", "vat_amount"),
                       ("مبلغ قابل پرداخت", "final_price")):
        ws.append([label, round(float(totals.get(key, 0)) / scale, 2)])
        ws.cell(row=ws.max_row, column=2).number_format = "#,##0"
        ws.cell(row=ws.max_row, column=1).font = Font(name="Vazirmatn", bold=(key == "final_price"), size=11)
        ws.cell(row=ws.max_row, column=2).font = Font(name="Vazirmatn", bold=(key == "final_price"), size=11)
    ws.append([])
    ws.append(["مبلغ به حروف: " + fa.money_to_words(totals.get("final_price", 0), cur, scale)])
    _autosize(ws, [8, 34, 54, 10, 10, 16, 18])
    wb.save(str(path))
    return path


# ---------------------------------------------------------------------------
# ۴) برش‌نامه و لیست خرید کارگاه
# ---------------------------------------------------------------------------

def export_workshop_xlsx(pack: Dict[str, Any], path: Path,
                         project: Optional[Dict[str, Any]] = None,
                         settings: Optional[Dict[str, str]] = None) -> Path:
    """
    خروجی اکسلِ «بسته‌ی کارگاهی»: سه برگه
      ۱. برش‌نامه      — تک‌تک قطعات با ابعاد میلی‌متری، برای ارّه و CNC
      ۲. خلاصه‌ی ورق   — چند ورق از هر متریال باید خرید/از انبار برداشت
      ۳. لیست خرید     — همه‌ی متریال و یراق، گروه‌بندی‌شده

    اگر `openpyxl` نصب نباشد، به‌جایش یک فایل CSV از برش‌نامه نوشته می‌شود تا
    کاربر دستِ خالی نماند.
    """
    proj = project or {}
    fa_digits = str((settings or {}).get("persian_digits", "1")) == "1"
    header = [f"{proj.get('title', '')}".strip() or "پروژه",
              f"کد: {proj.get('code', '')}".strip(),
              f"مشتری: {proj.get('customer', '')}".strip(),
              f"تاریخ: {fa.jdate_str()}"]
    header = [h for h in header if h and h not in ("کد:", "مشتری:", "تاریخ:")]

    parts = pack.get("cut_parts") or []
    sheets = pack.get("sheet_summary") or []
    groups = pack.get("shopping_groups") or []

    if not HAS_OPENPYXL:                      # خروجی جایگزین: CSV
        csv_path = path.with_suffix(".csv")
        with open(csv_path, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(["برش‌نامه"] + header)
            w.writerow(["ردیف", "بخش", "یونیت", "قطعه", "تعداد",
                        "طول (mm)", "عرض (mm)", "ضخامت (mm)", "متریال",
                        "نوار لبه", "مساحت (م²)"])
            for i, p in enumerate(parts, 1):
                w.writerow([i, p.get("section", ""), p.get("unit", ""), p.get("part", ""),
                            p.get("qty", 0), p.get("length_mm", 0), p.get("width_mm", 0),
                            p.get("thickness_mm", 0), p.get("material", ""),
                            "بله" if p.get("banding") else "خیر", p.get("area_sqm", 0)])
        return csv_path

    wb = Workbook()

    # --- برگه‌ی ۱: برش‌نامه ---
    ws = wb.active
    ws.title = "برش‌نامه"
    _rtl(ws)
    if header:
        ws.append([" | ".join(header)])
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=11)
        ws.cell(row=1, column=1).font = Font(name="Vazirmatn", bold=True, size=12)
    ws.append(["ردیف", "بخش", "یونیت", "قطعه", "تعداد", "طول (mm)", "عرض (mm)",
               "ضخامت", "متریال / ورق", "نوار لبه", "مساحت (م²)"])
    _style_header(ws, ws.max_row, 11)
    for i, p in enumerate(parts, 1):
        ws.append([
            fa.to_fa_digits(str(i)) if fa_digits else i,
            p.get("section", ""), p.get("unit", ""), p.get("part", ""),
            p.get("qty", 0),
            round(p.get("length_mm", 0)), round(p.get("width_mm", 0)),
            f"{p.get('thickness_mm', 0):g} میل",
            p.get("material", ""),
            ("دور تا دور" if p.get("banding_sides") == "دور تا دور"
             else p.get("banding_sides", "")) if p.get("banding") else "—",
            round(p.get("area_sqm", 0), 3),
        ])
    st = pack.get("stats") or {}
    ws.append([])
    ws.append(["جمع", "", "", f"{st.get('part_kinds', 0)} نوع قطعه",
               st.get("total_pieces", 0), "", "", "", "",
               "", st.get("total_area_sqm", 0)])
    for c in range(1, 12):
        ws.cell(row=ws.max_row, column=c).font = Font(name="Vazirmatn", bold=True)
    _autosize(ws, [6, 16, 20, 26, 8, 11, 11, 9, 30, 22, 12])
    ws.freeze_panes = "A3"

    # --- برگه‌ی ۲: خلاصه‌ی ورق ---
    ws2 = wb.create_sheet("خلاصه ورق")
    _rtl(ws2)
    ws2.append(["متریال / ورق", "دسته", "ضخامت", "ابعاد ورق (سم)",
                "مساحت قطعات (م²)", "تعداد قطعه", "ورق لازم", "بهای هر ورق", "جمع"])
    _style_header(ws2, 1, 9)
    cat_labels = {"body_sheet": "بدنه", "back_sheet": "فیبر پشت",
                  "door_sheet": "در", "drawer_sheet": "جعبه کشو"}
    for r in sheets:
        ws2.append([
            r.get("material", ""), cat_labels.get(r.get("cat"), r.get("cat", "")),
            f"{r.get('thickness_mm', 0):g} میل", r.get("sheet_size", ""),
            r.get("area_sqm", 0), r.get("pieces", 0), r.get("sheets", 0),
            r.get("price", 0), r.get("cost", 0),
        ])
    _autosize(ws2, [34, 12, 10, 16, 16, 12, 11, 15, 16])

    # --- برگه‌ی ۳: لیست خرید ---
    ws3 = wb.create_sheet("لیست خرید")
    _rtl(ws3)
    ws3.append(["#", "متریال / خدمت", "مقدار", "واحد", "برآورد بها", "خرید شد؟"])
    _style_header(ws3, 1, 6)
    idx = 1
    for g in groups:
        ws3.append([g.get("label", "")])
        cell = ws3.cell(row=ws3.max_row, column=1)
        cell.font = Font(name="Vazirmatn", bold=True, size=11)
        cell.fill = PatternFill("solid", fgColor=SOFT_HEX)
        ws3.merge_cells(start_row=ws3.max_row, start_column=1,
                        end_row=ws3.max_row, end_column=6)
        for r in g.get("rows", []):
            ws3.append([fa.to_fa_digits(str(idx)) if fa_digits else idx,
                        r.get("name", ""), round(r.get("qty", 0), 2),
                        r.get("unit_label", ""), round(r.get("amount", 0)), ""])
            idx += 1
        ws3.append(["", "جمع این گروه", "", "", round(g.get("total", 0)), ""])
        for c in range(1, 7):
            ws3.cell(row=ws3.max_row, column=c).font = Font(name="Vazirmatn", bold=True)
    _autosize(ws3, [5, 40, 12, 10, 18, 12])

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path
